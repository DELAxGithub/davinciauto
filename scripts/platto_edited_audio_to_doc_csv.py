#!/usr/bin/env python3
"""Fill a Platto Google Doc review table with edited-audio timecodes.

Inputs:
- DOCX export of the current Google Doc review table
- Whisper JSON for the edited master audio

Output:
- CSV with the same Doc table columns, adding/filling 編集後タイムコード
- review CSV with matching confidence and diagnostic notes
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any


HEADERS = ["カット候補", "Speaker Name", "素材イン点&ナレーション", "編集後タイムコード", "文字起こし"]
TIME_RE = re.compile(r"^\d{2}:\d{2}:\d{2}:\d{2}$")


@dataclass
class DocRow:
    row_number: int
    cut_candidate: str
    speaker: str
    source_or_na: str
    edited_timecode: str
    transcript: str


@dataclass
class MatchResult:
    row: DocRow
    edited_timecode: str
    match_status: str
    confidence: float
    segment_start: float | None
    segment_end: float | None
    matched_text: str
    note: str


def normalize_text(text: str) -> str:
    text = (text or "").strip()
    text = re.sub(r"\s+", "", text)
    drop = "、。,.!！?？「」『』（）()…ー〜-・ 　"
    return "".join(ch for ch in text if ch not in drop)


def seconds_to_tc(seconds: float, fps: float) -> str:
    frame = int(round(seconds * fps))
    fps_i = int(round(fps))
    hh = frame // (fps_i * 3600)
    frame %= fps_i * 3600
    mm = frame // (fps_i * 60)
    frame %= fps_i * 60
    ss = frame // fps_i
    ff = frame % fps_i
    return f"{hh:02d}:{mm:02d}:{ss:02d}:{ff:02d}"


def cell_text(cell: Any) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def load_doc_rows(docx_path: Path) -> list[DocRow]:
    from docx import Document

    document = Document(str(docx_path))
    if not document.tables:
        raise SystemExit(f"No tables found in {docx_path}")

    table = max(document.tables, key=lambda t: len(t.rows))
    rows: list[DocRow] = []
    for index, table_row in enumerate(table.rows):
        values = [cell_text(cell) for cell in table_row.cells]
        if values[: len(HEADERS)] == HEADERS:
            continue
        if len(values) < 5:
            continue
        if not any(value.strip() for value in values):
            continue
        rows.append(
            DocRow(
                row_number=index + 1,
                cut_candidate=values[0].strip(),
                speaker=values[1].strip(),
                source_or_na=values[2].strip(),
                edited_timecode=values[3].strip(),
                transcript=values[4].strip(),
            )
        )
    return rows


def load_segments(whisper_json: Path) -> list[dict[str, Any]]:
    data = json.loads(whisper_json.read_text(encoding="utf-8"))
    segments = []
    for segment in data.get("segments", []):
        text = (segment.get("text") or "").strip()
        if not text:
            continue
        segments.append(
            {
                "start": float(segment["start"]),
                "end": float(segment["end"]),
                "text": text,
                "norm": normalize_text(text),
            }
        )
    return segments


def build_doc_char_stream(rows: list[DocRow]) -> tuple[str, list[tuple[int, int]]]:
    stream_parts: list[str] = []
    row_ranges: list[tuple[int, int]] = []
    for row in rows:
        text = row.transcript or (row.source_or_na if row.speaker.upper().startswith("NA") else "")
        norm = normalize_text(text)
        start = sum(len(part) for part in stream_parts)
        stream_parts.append(norm)
        end = start + len(norm)
        row_ranges.append((start, end))
    return "".join(stream_parts), row_ranges


def build_segment_char_stream(segments: list[dict[str, Any]]) -> tuple[str, list[float]]:
    chars: list[str] = []
    char_times: list[float] = []
    for segment in segments:
        norm = segment["norm"]
        if not norm:
            continue
        duration = max(0.001, float(segment["end"]) - float(segment["start"]))
        for offset, char in enumerate(norm):
            chars.append(char)
            char_times.append(float(segment["start"]) + duration * (offset / max(len(norm), 1)))
    return "".join(chars), char_times


def build_sequence_matches(rows: list[DocRow], segments: list[dict[str, Any]], fps: float) -> list[MatchResult]:
    doc_stream, row_ranges = build_doc_char_stream(rows)
    edited_stream, edited_char_times = build_segment_char_stream(segments)
    matcher = SequenceMatcher(None, doc_stream, edited_stream, autojunk=False)

    row_hits: list[list[int]] = [[] for _ in rows]
    current_row = 0
    for block in matcher.get_matching_blocks():
        if block.size <= 0:
            continue
        doc_start = block.a
        doc_end = block.a + block.size
        while current_row < len(row_ranges) and row_ranges[current_row][1] <= doc_start:
            current_row += 1
        row_index = current_row
        while row_index < len(row_ranges) and row_ranges[row_index][0] < doc_end:
            row_start, row_end = row_ranges[row_index]
            overlap_start = max(row_start, doc_start)
            overlap_end = min(row_end, doc_end)
            if overlap_end > overlap_start:
                for doc_pos in range(overlap_start, overlap_end):
                    row_hits[row_index].append(block.b + (doc_pos - block.a))
            row_index += 1

    results: list[MatchResult] = []
    for row, hits, (row_start, row_end) in zip(rows, row_hits, row_ranges):
        target_len = row_end - row_start
        unique_hits = sorted(set(hits))
        coverage = len(unique_hits) / max(target_len, 1)
        edited_tc = ""
        status = "unmatched"
        start_sec: float | None = None
        end_sec: float | None = None
        matched_text = ""
        note = "sequence alignment"

        min_coverage = 0.44
        if target_len <= 5:
            min_coverage = 0.80
        elif target_len <= 10:
            min_coverage = 0.62

        if unique_hits and coverage >= min_coverage:
            first = min(unique_hits)
            last = max(unique_hits)
            if first < len(edited_char_times):
                start_sec = edited_char_times[first]
                end_sec = edited_char_times[min(last, len(edited_char_times) - 1)]
                edited_tc = seconds_to_tc(start_sec, fps)
                status = "matched" if coverage >= 0.62 else "low_confidence"
                matched_text = edited_stream[first : min(last + 1, len(edited_stream))]
        else:
            note = "insufficient sequence-alignment coverage"

        results.append(
            MatchResult(
                row=row,
                edited_timecode=edited_tc,
                match_status=status,
                confidence=round(coverage, 4),
                segment_start=start_sec,
                segment_end=end_sec,
                matched_text=matched_text,
                note=note,
            )
        )
    return results


def best_match(
    target: str,
    segments: list[dict[str, Any]],
    cursor: int,
    lookahead: int,
) -> tuple[int | None, float, str]:
    target_norm = normalize_text(target)
    if not target_norm:
        return None, 0.0, "empty target"

    best_index: int | None = None
    best_score = 0.0
    best_note = ""
    end = min(len(segments), cursor + lookahead)

    for index in range(cursor, end):
        combined = ""
        combined_text = []
        for span in range(0, 4):
            if index + span >= len(segments):
                break
            combined += segments[index + span]["norm"]
            combined_text.append(segments[index + span]["text"])
            if not combined:
                continue
            score = SequenceMatcher(None, target_norm, combined).ratio()
            if target_norm in combined or combined in target_norm:
                score = max(score, min(len(target_norm), len(combined)) / max(len(target_norm), len(combined)))
            if score > best_score:
                best_index = index
                best_score = score
                best_note = " + ".join(combined_text)
    return best_index, best_score, best_note


def build_matches(rows: list[DocRow], segments: list[dict[str, Any]], fps: float) -> list[MatchResult]:
    matches: list[MatchResult] = []
    cursor = 0

    for row in rows:
        target = row.transcript or (row.source_or_na if row.speaker.upper().startswith("NA") else "")
        index, score, matched_text = best_match(target, segments, cursor, lookahead=90)

        status = "unmatched"
        edited_tc = ""
        start: float | None = None
        end: float | None = None
        note = ""

        if index is not None and score >= 0.46:
            segment = segments[index]
            start = float(segment["start"])
            end = float(segment["end"])
            edited_tc = seconds_to_tc(start, fps)
            status = "matched" if score >= 0.62 else "low_confidence"
            cursor = max(cursor, index)
            if score >= 0.62:
                cursor = index + 1
            note = "text match"
        else:
            note = "no close edited-audio segment in forward window"

        matches.append(
            MatchResult(
                row=row,
                edited_timecode=edited_tc,
                match_status=status,
                confidence=round(score, 4),
                segment_start=start,
                segment_end=end,
                matched_text=matched_text,
                note=note,
            )
        )
    return matches


def write_outputs(matches: list[MatchResult], out_csv: Path, review_csv: Path) -> None:
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    with out_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        for match in matches:
            row = match.row
            writer.writerow(
                {
                    "カット候補": row.cut_candidate,
                    "Speaker Name": row.speaker,
                    "素材イン点&ナレーション": row.source_or_na,
                    "編集後タイムコード": match.edited_timecode or row.edited_timecode,
                    "文字起こし": row.transcript,
                }
            )

    with review_csv.open("w", newline="", encoding="utf-8-sig") as handle:
        fieldnames = [
            "row_number",
            "Speaker Name",
            "素材イン点&ナレーション",
            "編集後タイムコード",
            "文字起こし",
            "match_status",
            "confidence",
            "segment_start_sec",
            "segment_end_sec",
            "matched_text",
            "note",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for match in matches:
            row = match.row
            writer.writerow(
                {
                    "row_number": row.row_number,
                    "Speaker Name": row.speaker,
                    "素材イン点&ナレーション": row.source_or_na,
                    "編集後タイムコード": match.edited_timecode,
                    "文字起こし": row.transcript,
                    "match_status": match.match_status,
                    "confidence": match.confidence,
                    "segment_start_sec": "" if match.segment_start is None else round(match.segment_start, 3),
                    "segment_end_sec": "" if match.segment_end is None else round(match.segment_end, 3),
                    "matched_text": match.matched_text,
                    "note": match.note,
                }
            )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--whisper-json", type=Path, required=True)
    parser.add_argument("--out-csv", type=Path, required=True)
    parser.add_argument("--review-csv", type=Path, required=True)
    parser.add_argument("--fps", type=float, default=24.0)
    args = parser.parse_args()

    rows = load_doc_rows(args.docx)
    segments = load_segments(args.whisper_json)
    matches = build_sequence_matches(rows, segments, args.fps)
    write_outputs(matches, args.out_csv, args.review_csv)

    counts: dict[str, int] = {}
    for match in matches:
        counts[match.match_status] = counts.get(match.match_status, 0) + 1
    print(
        json.dumps(
            {
                "doc_rows": len(rows),
                "whisper_segments": len(segments),
                "counts": counts,
                "out_csv": str(args.out_csv),
                "review_csv": str(args.review_csv),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
