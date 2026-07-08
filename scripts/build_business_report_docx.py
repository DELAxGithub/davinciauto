import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT_DIR = Path("output/business_reports_2026_06")
DRIVE_OUTPUT_DIR = Path(
    "/Users/delaxpro/Library/CloudStorage/GoogleDrive-h.kodera@gmail.com/"
    "マイドライブ/Delaxプラッと/07_業務実績報告書"
)


REPORTS = [
    {
        # Final program titles must be checked against the official NHK episode list.
        # https://www.nhk.jp/p/rs/MPZ6XPWMV5/list/
        "filename": "【報告書】#37_カタツムリのしっぽはどこ?@千代田区.docx",
        "title": "NHK「プラッと」#37 カタツムリのしっぽはどこ?@千代田区",
        "contributors": ["東島沙弥佳氏", "千葉聡氏"],
        "shooting": "2026年4月20日 @皇居",
        "delivery": "2026年6月7日",
        "broadcast": "2026年6月16日 20:05~20:55 放送完了",
    },
    {
        "filename": "【報告書】#38_毒にも薬にもなる話@佃島.docx",
        "title": "NHK「プラッと」#38 毒にも薬にもなる話@佃島",
        "contributors": ["重田園江氏", "石塚真由美氏"],
        "shooting": "2026年5月31日 @佃島",
        "delivery": "2026年6月13日",
        "broadcast": "2026年6月23日 20:05~20:55 放送完了",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def add_para(cell, text, bold=False):
    p = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    run.font.size = Pt(10.5)
    run.bold = bold
    p.paragraph_format.space_after = Pt(0)
    return p


def add_table_row(table, label, value):
    row = table.add_row()
    set_cell_width(row.cells[0], 4.2)
    set_cell_width(row.cells[1], 11.5)
    set_cell_shading(row.cells[0], "F3F4F6")
    add_para(row.cells[0], label, bold=True)
    if isinstance(value, list):
        for i, item in enumerate(value):
            if i == 0:
                add_para(row.cells[1], item)
            else:
                row.cells[1].add_paragraph(item)
                p = row.cells[1].paragraphs[-1]
                p.runs[0].font.name = "Arial"
                p.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
                p.runs[0].font.size = Pt(10.5)
                p.paragraph_format.space_after = Pt(0)
    else:
        add_para(row.cells[1], value)


def build_report(data):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("業務実施報告書")
    r.bold = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    r.font.size = Pt(18)

    doc.add_paragraph()
    for line in [
        "株式会社NHKエンタープライズ 御中",
        "",
        "〒150-0047",
        "東京都渋谷区神山町4-14",
        "第1制作センター 文化部",
        "丸山 俊一様",
        "下記のとおり、業務を完了いたしましたので、ご報告申し上げます。",
    ]:
        doc.add_paragraph(line)

    center = doc.add_paragraph("記")
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    add_table_row(table, "1. 契約名称", "放送番組制作業務委託(演出)契約書")
    add_table_row(table, "2. 契約期間", "2026年4月1日 ~ 2026年6月30日")
    add_table_row(table, "3. 番組名", data["title"])
    add_table_row(
        table,
        "4. 委託業務内容",
        [
            "甲の制作統括の下での、本件番組の演出業務および権利処理を含むこれに付随する業務",
            "・構成、演出、取材、編集等の業務",
            "・主な番組寄与者を除く権利処理",
            "・上記項目の業務遂行のために必要な本番組の打合せ参加",
        ],
    )
    add_table_row(table, "5. 業務実施期間", "2026年4月1日 ~ 2026年6月30日")
    add_table_row(table, "6. 主な番組寄与者(出演者)", data["contributors"])
    add_table_row(table, "7. 撮影実施日", data["shooting"])
    add_table_row(
        table,
        "8. 成果物および納品内容",
        [
            "(1) ラジオ版",
            "・尺: 50分",
            "・フォーマット: WAV",
            f"・納品日: {data['delivery']}",
            "・納品方法: BOX",
        ],
    )
    add_table_row(table, "9. 放送実績", f"(1) ラジオ第1: {data['broadcast']}")
    add_table_row(table, "10. 委託費", ["金500,000円(税別)", "消費税 50,000円", "合計 金550,000円"])
    add_table_row(table, "11. 支払期日", "すべての委託業務完了日の翌月末日")

    doc.add_paragraph()
    end = doc.add_paragraph("以上")
    end.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for line in [
        "2026年6月30日",
        "石川県金沢市桂町ロ239番地",
        "株式会社DELAX",
        "代表取締役 小寺 寛志",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    out = OUT_DIR / data["filename"]
    doc.save(out)
    return out


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    DRIVE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for report in REPORTS:
        out = build_report(report)
        drive_out = DRIVE_OUTPUT_DIR / out.name
        shutil.copy2(out, drive_out)
        print(out)
        print(drive_out)


if __name__ == "__main__":
    main()
